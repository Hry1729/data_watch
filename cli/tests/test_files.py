from __future__ import annotations

import json

import pytest
from openpyxl import Workbook

from cli.errors import CliError, EXIT_USAGE
from cli.files import (
    load_json_object,
    load_workbook_rows,
    parse_ids,
    require_confirmation,
)


def test_device_workbook_uses_second_sheet_and_header_aliases(tmp_path) -> None:
    path = tmp_path / "devices.xlsx"
    workbook = Workbook()
    workbook.active.title = "说明"
    sheet = workbook.create_sheet("设备数据")
    sheet.append(["设备编号", "deviceName", "关联驱动", "设备状态"])
    sheet.append([" PLC001 ", "Main PLC", "S7", "enabled"])
    sheet.append(["", "", "", ""])
    workbook.save(path)

    rows = load_workbook_rows(str(path), "device")

    assert rows == [
        {
            "deviceCode": "PLC001",
            "deviceName": "Main PLC",
            "driverName": "S7",
            "deviceStatus": "enabled",
        }
    ]


def test_json_payload_requires_an_object(tmp_path) -> None:
    path = tmp_path / "payload.json"
    path.write_text(json.dumps([1, 2, 3]), encoding="utf-8")

    with pytest.raises(CliError) as captured:
        load_json_object(str(path))

    assert captured.value.exit_code == EXIT_USAGE


def test_confirmation_and_batch_ids_are_strict() -> None:
    require_confirmation("PLC001", "PLC001")
    assert parse_ids("1, 2,3") == [1, 2, 3]

    with pytest.raises(CliError):
        require_confirmation("plc001", "PLC001")
    with pytest.raises(CliError):
        parse_ids("1,1")


def test_json_pipe_decodes_utf8_independently_of_console_encoding(monkeypatch):
    import io
    import json
    from cli.files import load_json_object
    payload = {"deviceName": "设备测试"}
    stream = io.TextIOWrapper(io.BytesIO(json.dumps(payload, ensure_ascii=False).encode("utf-8")), encoding="gbk")
    monkeypatch.setattr("cli.files.sys.stdin", stream)
    assert load_json_object("-") == payload


def test_point_workbook_preserves_false_and_zero(tmp_path) -> None:
    """B050：布尔 FALSE / 数值 0 不能被 truthiness 抹成空字符串（否则后端按缺省 true 处理）。"""
    path = tmp_path / "points.xlsx"
    workbook = Workbook()
    workbook.active.title = "说明"
    sheet = workbook.create_sheet("点位数据")
    sheet.append(["设备编号", "点位名称", "是否只读", "是否存储"])
    sheet.append(["PLC001", "p1", False, 0])
    sheet.append(["PLC001", "p2", True, 1])
    sheet.append(["PLC001", "p3", None, None])
    workbook.save(path)

    rows = load_workbook_rows(str(path), "point")

    assert [(row["isReadonly"], row["isStorage"]) for row in rows] == [
        (False, 0),
        (True, 1),
        ("", ""),
    ]
