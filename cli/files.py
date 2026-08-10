from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from cli.errors import CliError, EXIT_API, EXIT_USAGE


DEVICE_FIELDS = {
    "设备类型": "deviceType",
    "设备编号": "deviceCode",
    "设备名称": "deviceName",
    "关联驱动": "driverName",
    "设备状态": "deviceStatus",
    "描述": "description",
}
POINT_FIELDS = {
    "设备编号": "deviceCode",
    "点位序号": "varSeq",
    "点位名称": "pointName",
    "IO地址": "ioAddress",
    "数据类型": "dataType",
    "驱动名称": "driverName",
    "是否只读": "isReadonly",
    "描述": "description",
    "采集周期(ms)": "scanRate",
    "是否存储": "isStorage",
}


def load_json_object(path_value: str) -> dict[str, Any]:
    path = _input_file(path_value)
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as exc:
        raise CliError("file_error", f"无法读取 JSON 文件: {path}", EXIT_USAGE) from exc
    if not isinstance(payload, dict):
        raise CliError("file_error", "JSON 文件根节点必须是对象", EXIT_USAGE)
    return payload


def load_workbook_rows(path_value: str, kind: str) -> list[dict[str, Any]]:
    path = _input_file(path_value)
    if path.suffix.lower() != ".xlsx":
        raise CliError("file_error", "导入文件必须是 .xlsx", EXIT_USAGE)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise CliError("dependency_error", "缺少 openpyxl 依赖", EXIT_API) from exc

    fields = DEVICE_FIELDS if kind == "device" else POINT_FIELDS
    aliases = {
        _normalize_header(alias): field
        for label, field in fields.items()
        for alias in (label, field)
    }
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.worksheets[1] if len(workbook.worksheets) > 1 else workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise CliError("file_error", "工作簿没有表头", EXIT_USAGE)
        mapped_headers = [aliases.get(_normalize_header(value)) for value in headers]
        if not any(mapped_headers):
            raise CliError("file_error", "工作簿表头与导入模板不匹配", EXIT_USAGE)
        result: list[dict[str, Any]] = []
        for values in rows:
            item: dict[str, Any] = {}
            for field, value in zip(mapped_headers, values, strict=False):
                if not field:
                    continue
                item[field] = value.strip() if isinstance(value, str) else (value or "")
            if any(str(value).strip() for value in item.values()):
                result.append(item)
        workbook.close()
    except CliError:
        raise
    except Exception as exc:
        raise CliError("file_error", f"无法解析 XLSX 文件: {path}", EXIT_USAGE) from exc
    if not result:
        raise CliError("file_error", "工作簿没有可导入的数据行", EXIT_USAGE)
    return result


def write_download(
    output_value: str,
    content: bytes,
    server_filename: str,
) -> dict[str, Any]:
    output = Path(output_value).expanduser()
    if output.exists() and output.is_dir():
        output = output / server_filename
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    try:
        temporary.write_bytes(content)
        temporary.replace(output)
    except OSError as exc:
        raise CliError("file_error", f"无法写入文件: {output}", EXIT_API) from exc
    return {
        "output": str(output.resolve()),
        "fileName": server_filename,
        "bytes": len(content),
    }


def require_confirmation(value: str, expected: str) -> None:
    if value != expected:
        raise CliError(
            "confirmation_required",
            f"--confirm 必须精确等于 {expected}",
            EXIT_USAGE,
            details={"expected": expected},
        )


def parse_ids(value: str) -> list[int]:
    try:
        ids = [int(item.strip()) for item in value.split(",") if item.strip()]
    except ValueError as exc:
        raise CliError("usage_error", "--ids 必须是逗号分隔的整数", EXIT_USAGE) from exc
    if not ids or any(item <= 0 for item in ids) or len(set(ids)) != len(ids):
        raise CliError("usage_error", "--ids 必须包含不重复的正整数", EXIT_USAGE)
    return ids


def _input_file(path_value: str) -> Path:
    path = Path(path_value).expanduser()
    if not path.is_file():
        raise CliError("file_error", f"文件不存在: {path}", EXIT_USAGE)
    return path


def _normalize_header(value: Any) -> str:
    normalized = re.sub(r"[()（）]", "", str(value or "").strip())
    return re.sub(r"[\s_-]+", "", normalized).lower()
