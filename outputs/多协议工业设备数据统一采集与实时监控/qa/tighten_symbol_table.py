from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

docx = Path(r"D:\python\device_alarm\data_watch\outputs\多协议工业设备数据统一采集与实时监控\一种多协议工业设备数据的统一采集与实时监控方法及系统_20260714180446.docx")
doc = Document(docx)
table = doc.tables[2]
for row in table.rows:
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for tag in ("top", "bottom"):
            node = tc_mar.find(qn(f"w:{tag}"))
            if node is None:
                node = OxmlElement(f"w:{tag}")
                tc_mar.append(node)
            node.set(qn("w:w"), "35")
            node.set(qn("w:type"), "dxa")
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
doc.save(docx)
print(docx)
