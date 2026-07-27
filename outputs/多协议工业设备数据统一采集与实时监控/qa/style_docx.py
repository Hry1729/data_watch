from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\python\device_alarm\data_watch")
DOCX = ROOT / "outputs" / "多协议工业设备数据统一采集与实时监控" / "一种多协议工业设备数据的统一采集与实时监控方法及系统_20260714180446.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_FILL = "F4F6F9"
MUTED = "666666"


def set_east_asia(run, east_asia="宋体", ascii_font="Calibri"):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_east_asia(run)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    tail = paragraph.add_run(" 页")
    set_east_asia(tail)


INLINE_REPLACEMENTS = {
    r"\(S_k\)": "Sₖ",
    r"\(n\)": "n",
    r"\(\hat{s}_{i,n-1}\)": "ŝᵢ,ₙ₋₁",
    r"\(R_n\)": "Rₙ",
    r"\(L_{\max}\)": "Lₘₐₓ",
    r"\(\tau_u\)": "τᵤ",
    r"\(N\)": "N",
    r"\(t_k-t_{k-1}\)": "tₖ−tₖ₋₁",
    r"\(i\)": "i",
    r"\(k\)": "k",
    r"\(s_{i,k}\)": "sᵢ,ₖ",
    r"\(S_k\)": "Sₖ",
    r"\(\Delta_k\)": "Δₖ",
    r"\(t_k\)": "tₖ",
    r"\(B_n\)": "Bₙ",
    r"\(\hat{s}_{i,n}\)": "ŝᵢ,ₙ",
    r"\(k^*\)": "k*",
    r"\(L_{\max}\)": "Lₘₐₓ",
}


def readable_math(text):
    if "\\Delta_k" in text and "tag{1}" in text:
        return "Δₖ = { sᵢ,ₖ | sᵢ,ₖ ≠ sᵢ,ₖ₋₁，i = 1，…，N }    （1）"
    if "\\bigcup" in text and "tag{2}" in text:
        return "Bₙ = ⋃ Δₖ，其中（n−1）τᵤ < tₖ ≤ nτᵤ    （2）"
    if "\\hat{s}_{i,n}" in text and "tag{3}" in text:
        return "ŝᵢ,ₙ = sᵢ,ₖ*，其中 k* = max { k | sᵢ,ₖ ∈ Bₙ }    （3）"
    if "|R|" in text and "tag{4}" in text:
        return "0 ≤ |R| ≤ Lₘₐₓ    （4）"
    result = text
    for old, new in INLINE_REPLACEMENTS.items():
        result = result.replace(old, new)
    result = result.replace(r"\(", "").replace(r"\)", "")
    return result


def replace_paragraph_text(paragraph, text, math=False):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    set_east_asia(run, east_asia="Cambria Math" if math else "宋体", ascii_font="Cambria Math" if math else "Calibri")
    run.font.size = Pt(10.5 if math else 11)
    return run


def walk_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from walk_table_paragraphs(nested)


doc = Document(DOCX)

# narrative_proposal preset, with one named override: Chinese glyphs use SimSun;
# headings use Microsoft YaHei for consistent CJK rendering.
for section in doc.sections:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.333

heading_tokens = {
    "Heading 1": (16, BLUE, 18, 10),
    "Heading 2": (16, BLUE, 18, 10),
    "Heading 3": (13, BLUE, 12, 6),
    "Heading 4": (12, DARK_BLUE, 8, 4),
}
for style_name, (size, color, before, after) in heading_tokens.items():
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Number"):
    if style_name in doc.styles:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

all_paragraphs = list(doc.paragraphs)
for table in doc.tables:
    all_paragraphs.extend(walk_table_paragraphs(table))

for paragraph in all_paragraphs:
    original = paragraph.text
    is_block_math = "tag{" in original and ("$$" in original or r"\[" in original)
    converted = readable_math(original)
    if converted != original:
        replace_paragraph_text(paragraph, converted, math=is_block_math)
    for run in paragraph.runs:
        if is_block_math:
            set_east_asia(run, east_asia="Cambria Math", ascii_font="Cambria Math")
            run.font.size = Pt(10.5)
        elif paragraph.style.name.startswith("Heading"):
            set_east_asia(run, east_asia="微软雅黑", ascii_font="Calibri")
        else:
            set_east_asia(run)
    if is_block_math:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.keep_together = True
    elif paragraph.style.name == "Normal":
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.333
    if paragraph.style.name.startswith("Heading"):
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
    if paragraph._p.xpath(".//w:drawing"):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(10)

# Title is a named first-paragraph override to the narrative preset.
title = doc.paragraphs[0]
title.style = doc.styles["Normal"]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(14)
for run in title.runs:
    set_east_asia(run, east_asia="微软雅黑", ascii_font="Calibri")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

for index in range(1, min(8, len(doc.paragraphs))):
    paragraph = doc.paragraphs[index]
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15

# Full-width fixed-DXA tables: 6.5 in = 9360 DXA, tblInd = 120 DXA.
for table in doc.tables:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    cols = len(table.columns)
    if cols == 5:
        widths = [1440, 2376, 1224, 2448, 1872]
    elif cols == 4:
        widths = [1728, 1584, 2160, 3888]
    elif cols == 3:
        widths = [1728, 4320, 3312]
    else:
        widths = [9360 // cols] * cols
        widths[-1] += 9360 - sum(widths)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(col_index, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade_cell(cell, LIGHT_FILL)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.08
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_east_asia(run)
                    run.font.size = Pt(9 if cols >= 4 else 9.5)
                    if row_index == 0:
                        run.font.bold = True
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)

# Quiet running header and page-number footer.
for section in doc.sections:
    header = section.header
    for paragraph in header.paragraphs:
        paragraph.text = ""
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    left = hp.add_run("技术交底书")
    set_east_asia(left, east_asia="微软雅黑", ascii_font="Calibri")
    left.font.size = Pt(9)
    left.font.bold = True
    left.font.color.rgb = RGBColor.from_string(MUTED)
    right = hp.add_run("\t多协议工业设备数据统一采集与实时监控")
    set_east_asia(right, east_asia="微软雅黑", ascii_font="Calibri")
    right.font.size = Pt(9)
    right.font.color.rgb = RGBColor.from_string(MUTED)
    tabs = hp.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5))

    footer = section.footer
    for paragraph in footer.paragraphs:
        paragraph.text = ""
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_field(fp)
    for run in fp.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)

doc.core_properties.title = "一种多协议工业设备数据的统一采集与实时监控方法及系统"
doc.core_properties.subject = "发明专利技术交底书"
doc.core_properties.keywords = "多协议；工业设备；统一采集；实时监控；差分推送；列式流式处理"
doc.save(DOCX)
print(DOCX)
