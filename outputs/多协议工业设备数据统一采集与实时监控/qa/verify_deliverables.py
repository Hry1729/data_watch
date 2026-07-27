from pathlib import Path
from zipfile import ZipFile

from docx import Document

root = Path(r"D:\python\device_alarm\data_watch\outputs\多协议工业设备数据统一采集与实时监控")
stem = "一种多协议工业设备数据的统一采集与实时监控方法及系统_20260714180446"
md_path = root / f"{stem}.md"
docx_path = root / f"{stem}.docx"
pdf_path = root / "qa" / "rendered" / "patent-disclosure.pdf"
page_images = sorted((root / "qa" / "rendered").glob("page-*.png"))

assert md_path.is_file() and md_path.stat().st_size > 20_000
assert docx_path.is_file() and docx_path.stat().st_size > 100_000
assert pdf_path.is_file() and pdf_path.stat().st_size > 100_000
assert len(page_images) == 13

md = md_path.read_text(encoding="utf-8")
title = "一种多协议工业设备数据的统一采集与实时监控方法及系统"
assert f"**案件名称**：{title}" in md
assert md.count("```mermaid") == 2
assert md.count("<!-- ![图示") == 2
assert md.count("https://patents.google.com/patent/") == 5
for chapter in ("## 一、", "## 二、", "## 三、", "## 四、", "## 五、", "## 六、"):
    assert chapter in md
for forbidden in (
    "patent-disclosure-skill",
    "disclosure_draft",
    "教学示例",
    "不构成法律",
    "[待填写]",
    "Perspective",
):
    assert forbidden not in md
assert "$$" not in md
assert "\\[" in md and "\\(" in md

with ZipFile(docx_path) as archive:
    assert archive.testzip() is None
    assert "word/document.xml" in archive.namelist()

doc = Document(docx_path)
assert doc.core_properties.title == title
assert len(doc.tables) == 4
assert len(doc.inline_shapes) == 2
assert len(doc.sections) == 1
section = doc.sections[0]
assert round(section.page_width.inches, 2) == 8.5
assert round(section.page_height.inches, 2) == 11.0
assert round(section.left_margin.inches, 2) == 1.0
assert round(section.right_margin.inches, 2) == 1.0

texts = [paragraph.text for paragraph in doc.paragraphs]
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            texts.extend(paragraph.text for paragraph in cell.paragraphs)
joined = "\n".join(texts)
assert title in joined
assert "$$" not in joined and "\\[" not in joined and "\\(" not in joined
assert "Δₖ" in joined and "τᵤ" in joined and "Lₘₐₓ" in joined
assert "技术交底书" in section.header.paragraphs[0].text
assert "第 " in section.footer.paragraphs[0].text

print(f"MD_OK bytes={md_path.stat().st_size} mermaid=2 urls=5 chapters=6")
print(f"DOCX_OK bytes={docx_path.stat().st_size} tables=4 images=2 page=Letter margins=1in")
print(f"RENDER_OK pdf_bytes={pdf_path.stat().st_size} pages={len(page_images)}")
