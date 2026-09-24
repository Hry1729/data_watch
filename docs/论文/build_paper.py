from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "DataWatch_工业数据底座论文.md"
OUTPUT = ROOT / "DataWatch_工业数据底座论文.docx"
FIGURES = {
    "figures/datawatch_system_architecture.png": ROOT / "figures" / "datawatch_system_architecture.png",
    "figures/datawatch_block_coalescing.png": ROOT / "figures" / "datawatch_block_coalescing.png",
}

ACCENT = "176B73"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(86, 98, 110)
LIGHT = "EAF3F4"


def set_run_font(run, latin="Times New Roman", east_asia="宋体", size=10.5, bold=None, italic=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = TEXT
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, attrs in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in attrs.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_repeatable_heading(paragraph):
    set_keep_with_next(paragraph)
    paragraph.paragraph_format.keep_together = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=8.5)
    run.font.color.rgb = MUTED


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


CITATION_RE = re.compile(r"\[\d+(?:[–-]\d+)?(?:,\s*\d+)*\]")
TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*|\[\d+(?:[–-]\d+)?(?:,\s*\d+)*\]|https?://\S+)")


def add_inline(paragraph, text, *, size=10.5, color=TEXT):
    cursor = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=size)
            run.font.color.rgb = color
        token = match.group(0)
        if token.startswith("http"):
            url = token.rstrip(".,;，。")
            suffix = token[len(url) :]
            add_hyperlink(paragraph, url, url)
            if suffix:
                run = paragraph.add_run(suffix)
                set_run_font(run, size=size)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
            run.font.color.rgb = color
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, latin="Consolas", east_asia="等线", size=max(size - 0.7, 8.5))
            run.font.color.rgb = RGBColor(51, 65, 85)
        elif CITATION_RE.fullmatch(token):
            # 参考文献序号按中文期刊习惯以上标呈现
            run = paragraph.add_run(token)
            set_run_font(run, size=max(size - 3.2, 7))
            run.font.superscript = True
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
            run.font.color.rgb = color
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)
        run.font.color.rgb = color


def set_paragraph_body(paragraph, *, indent=True, compact=False):
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = 1.28
    fmt.space_after = Pt(3 if compact else 4.5)
    fmt.space_before = Pt(0)
    if indent:
        fmt.first_line_indent = Pt(21)
    fmt.widow_control = True


def add_body(document, text, *, indent=True, size=10.5, compact=False):
    paragraph = document.add_paragraph()
    set_paragraph_body(paragraph, indent=indent, compact=compact)
    add_inline(paragraph, text, size=size)
    return paragraph


def add_title(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, east_asia="黑体", size=18, bold=True)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    return paragraph


def add_english_title(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=11.5, bold=True)
    run.font.color.rgb = RGBColor(45, 55, 72)
    return paragraph


def add_author(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(13)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, east_asia="宋体", size=9.5)
    run.font.color.rgb = MUTED


def add_heading(document, text, level):
    paragraph = document.add_paragraph()
    set_repeatable_heading(paragraph)
    fmt = paragraph.paragraph_format
    if level == 2:
        fmt.space_before = Pt(13)
        fmt.space_after = Pt(5)
        run = paragraph.add_run(text)
        set_run_font(run, east_asia="黑体", size=13, bold=True)
        run.font.color.rgb = RGBColor.from_string(ACCENT)
        p_pr = paragraph._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "75A9AD")
        border.append(bottom)
        p_pr.append(border)
    else:
        fmt.space_before = Pt(9)
        fmt.space_after = Pt(3.5)
        run = paragraph.add_run(text)
        set_run_font(run, east_asia="黑体", size=11, bold=True)
    return paragraph


def add_list_item(document, text, ordered=False):
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(0.64)
    fmt.first_line_indent = Cm(-0.5)
    fmt.line_spacing = 1.22
    fmt.space_after = Pt(3)
    add_inline(paragraph, text, size=10.3)
    return paragraph


def add_equation(document, raw_lines):
    raw = " ".join(line.strip() for line in raw_lines)
    # 提取行尾的公式编号（（1）/（2）或 (1)/(2)），按中文期刊习惯右对齐排版
    number = ""
    match = re.search(r"[（(]\s*(\d+)\s*[）)]\s*$", raw)
    if match:
        number = f"（{match.group(1)}）"
        raw = raw[: match.start()].strip().rstrip(" ,，")
    if raw.startswith("x=") or raw.startswith("x ="):
        text = "x = (d, v, value, t, q),"
    elif raw.startswith("E=") or raw.startswith("E ="):
        text = "E = (device_code, device_type, group_code, scan_rate, ts, config_version, values, qualities)."
    else:
        text = raw.replace("\\_", "_").replace("\\text{-}", "-").replace("\\", "")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_together = True
    if number:
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(8.2), WD_TAB_ALIGNMENT.CENTER)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(16.4), WD_TAB_ALIGNMENT.RIGHT)
        run = paragraph.add_run("\t" + text + "\t" + number)
    else:
        run = paragraph.add_run(text)
    set_run_font(run, latin="Cambria Math", east_asia="Cambria Math", size=10.5, italic=True)


def table_widths(headers):
    key = tuple(headers)
    if len(headers) == 4:
        return [2.2, 4.0, 4.2, 5.6]
    if len(headers) == 3:
        return [3.2, 5.2, 7.6]
    if len(headers) == 7:
        return [3.2, 2.0, 1.8, 1.8, 1.9, 3.3, 2.5]
    return [16.0 / len(headers)] * len(headers)


def add_table(document, rows):
    headers = rows[0]
    table = document.add_table(rows=len(rows), cols=len(headers))
    table.autofit = False
    table.alignment = 1
    widths = table_widths(headers)
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Pt(18)
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Cm(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            paragraph.paragraph_format.keep_together = True
            add_inline(paragraph, value, size=8.5 if len(headers) >= 7 else 8.8)
            if row_index == 0:
                cell._tc.get_or_add_tcPr().append(OxmlElement("w:shd"))
                cell._tc.tcPr[-1].set(qn("w:fill"), LIGHT)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(29, 78, 82)
    set_repeat_table_header(table.rows[0])
    edge = {"val": "single", "sz": "10", "space": "0", "color": ACCENT}
    mid = {"val": "single", "sz": "6", "space": "0", "color": "75A9AD"}
    for cell in table.rows[0].cells:
        set_cell_border(cell, top=edge, bottom=mid)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom=edge)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def configure_document(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("DATAWATCH  ·  YANGSHAN PHASE IV INDUSTRIAL IOT DATA PLATFORM")
    set_run_font(run, size=7.5, bold=True)
    run.font.color.rgb = MUTED
    p_pr = header._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "B9D2D4")
    border.append(bottom)
    p_pr.append(border)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    core = document.core_properties
    core.title = "DataWatch：以洋山四期自动化码头为例的契约保持型工业物联网数据采集与监控平台"
    core.subject = "以洋山四期自动化码头为例的工业物联网数据采集与监控平台论文"
    core.keywords = "洋山四期自动化码头, 工业物联网, 异构设备采集, 多点位块合并, 插件化, 边缘中心协同, AI智能体, DataWatch"
    core.comments = "Generated from the repository-grounded manuscript; author metadata remains to be completed."


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    index = 0
    abstract_mode = False
    reference_mode = False
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped == "## 投稿前作者核对清单（定稿时删除）":
            break
        if stripped in {"", "---"}:
            index += 1
            continue
        if stripped.startswith("# "):
            add_title(document, stripped[2:])
        elif stripped.startswith("**DataWatch:"):
            add_english_title(document, stripped[2:-2])
        elif stripped.startswith("作者及单位："):
            add_author(document, stripped)
        elif stripped.startswith(("中图分类号", "文献标志码", "基金项目", "作者简介")):
            # 期刊前置信息行：左对齐、无首行缩进、字号略小
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, stripped, size=9)
        elif stripped.startswith("## "):
            heading = stripped[3:]
            add_heading(document, heading, 2)
            abstract_mode = heading in {"摘要", "Abstract"}
            reference_mode = heading == "参考文献"
        elif stripped.startswith("### "):
            add_heading(document, stripped[4:], 3)
            abstract_mode = False
        elif stripped.startswith("!["):
            image_match = re.search(r"\(([^)]+)\)", stripped)
            if image_match is None or image_match.group(1) not in FIGURES:
                raise ValueError(f"unknown manuscript figure: {stripped}")
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run()
            run.add_picture(str(FIGURES[image_match.group(1)]), width=Inches(6.35))
        elif stripped.startswith("**图 "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.right_indent = Cm(0.7)
            paragraph.paragraph_format.space_after = Pt(7)
            add_inline(paragraph, stripped, size=8.7, color=MUTED)
        elif stripped.startswith("**表 "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, stripped, size=9.2)
        elif stripped.startswith("|"):
            rows, next_index = parse_table(lines, index)
            add_table(document, rows)
            index = next_index - 1
        elif stripped == "\\[":
            equation_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != "\\]":
                equation_lines.append(lines[index])
                index += 1
            add_equation(document, equation_lines)
        elif re.match(r"^\d+\.\s", stripped):
            if reference_mode:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.65)
                paragraph.paragraph_format.first_line_indent = Cm(-0.65)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.08
                add_inline(paragraph, stripped, size=8.8)
            else:
                add_list_item(document, stripped, ordered=True)
        elif stripped.startswith("- "):
            add_list_item(document, "• " + stripped[2:])
        else:
            size = 9.4 if abstract_mode else 10.5
            indent = not abstract_mode and not stripped.startswith("**关键词") and not stripped.startswith("**Keywords")
            paragraph = add_body(document, stripped, indent=indent, size=size, compact=abstract_mode)
            if abstract_mode:
                paragraph.paragraph_format.left_indent = Cm(0.35)
                paragraph.paragraph_format.right_indent = Cm(0.35)

        index += 1

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
