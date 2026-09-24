from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

import pdfplumber
from docx import Document


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "DataWatch_工业数据底座论文.docx"
PDF = ROOT / "DataWatch_工业数据底座论文.pdf"


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    problems: list[str] = []
    document = Document(DOCX)
    paragraph_text = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    if len(document.tables) != 5:
        problems.append(f"expected 5 tables, found {len(document.tables)}")
    if any("投稿前作者核对清单" in text for text in paragraph_text):
        problems.append("author-only checklist leaked into DOCX")
    for expected in ("摘要", "Abstract", "1 引言", "6 工程建设过程与实现", "7 系统测试与工程观察", "9 结论", "参考文献"):
        if expected not in paragraph_text:
            problems.append(f"missing section: {expected}")

    source = (ROOT / "DataWatch_工业数据底座论文.md").read_text(encoding="utf-8")
    abstract_match = re.search(r"## 摘要\s+(.*?)\s+## Abstract", source, re.S)
    if not abstract_match:
        problems.append("unable to locate Chinese abstract in source")
    else:
        abstract_text = abstract_match.group(1)
        for forbidden in ("993.89", "1999.99", "73.44", "125.64", "300 MB"):
            if forbidden in abstract_text:
                problems.append(f"benchmark value leaked into abstract: {forbidden}")

    with zipfile.ZipFile(DOCX) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        if len(media) != 2:
            problems.append(f"expected 2 embedded figures, found {len(media)}")

    with pdfplumber.open(PDF) as pdf:
        print(f"pages={len(pdf.pages)} tables={len(document.tables)} media={len(media)}")
        if not 9 <= len(pdf.pages) <= 18:
            problems.append(f"unexpected page count: {len(pdf.pages)}")
        previous_last = ""
        for number, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(use_text_flow=True, keep_blank_chars=False)
            text = page.extract_text() or ""
            body_words = [w for w in words if 38 <= float(w["top"]) <= page.height - 55]
            top = min((float(w["top"]) for w in body_words), default=0)
            bottom = max((float(w["bottom"]) for w in body_words), default=0)
            images = len(page.images)
            tables = len(page.extract_tables())
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            first = next((line for line in lines if "DATAWATCH" not in line), "")
            last = next((line for line in reversed(lines) if not re.fullmatch(r"\d+", line)), "")
            print(
                f"page={number} chars={len(text)} body_top={top:.1f} "
                f"body_bottom={bottom:.1f} images={images} detected_tables={tables} "
                f"first={first[:55]!r} last={last[:55]!r}"
            )
            if len(text.strip()) < 120 and images == 0:
                problems.append(f"page {number} may be unexpectedly sparse")
            if bottom > page.height - 45:
                problems.append(f"page {number} has text too close to lower edge: {bottom:.1f}")
            if re.match(r"^(\d+(?:\.\d+)*\s|摘要$|Abstract$|参考文献$)", last):
                problems.append(f"page {number} appears to end with an orphan heading: {last}")
            previous_last = last

    if problems:
        print("QA_FAIL")
        for problem in problems:
            print(f"- {problem}")
        raise SystemExit(1)
    print("QA_OK")


if __name__ == "__main__":
    main()
